#!/usr/bin/env python3
"""Build a budgeting workbook from bank CSV exports and an Excel template."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
import re
from typing import Dict, Iterable, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows


DEFAULT_CATEGORY_MAP: Dict[str, str] = {
    "grocery|whole foods|trader joe": "Groceries",
    "rent|apartment|landlord": "Rent",
    "uber|lyft|metro|transit": "Transportation",
    "netflix|spotify|hulu": "Subscriptions",
    "electric|water|gas bill|utility": "Utilities",
    "restaurant|cafe|coffee|doordash|ubereats": "Dining",
    "payroll|salary|direct deposit": "Income",
}


@dataclass
class BuilderConfig:
    csv_path: Path
    template_path: Path
    output_path: Path
    date_col: str
    amount_col: str
    description_col: str
    debit_positive: bool


def parse_args() -> BuilderConfig:
    parser = argparse.ArgumentParser(
        description="Convert bank CSV data into a budgeting spreadsheet based on an Excel template."
    )
    parser.add_argument("--csv", required=True, type=Path, help="Path to bank transactions CSV")
    parser.add_argument(
        "--template", required=True, type=Path, help="Path to baseline Excel workbook"
    )
    parser.add_argument(
        "--output",
        required=True,
        type=Path,
        help="Path for generated budget workbook (.xlsx)",
    )
    parser.add_argument("--date-col", default="Date", help="CSV column containing transaction date")
    parser.add_argument(
        "--amount-col", default="Amount", help="CSV column containing transaction amount"
    )
    parser.add_argument(
        "--description-col",
        default="Description",
        help="CSV column containing transaction description",
    )
    parser.add_argument(
        "--debit-positive",
        action="store_true",
        help="Set if your CSV uses positive numbers for expenses and negative for income.",
    )

    args = parser.parse_args()
    return BuilderConfig(
        csv_path=args.csv,
        template_path=args.template,
        output_path=args.output,
        date_col=args.date_col,
        amount_col=args.amount_col,
        description_col=args.description_col,
        debit_positive=args.debit_positive,
    )


def categorize(description: str, mapping: Dict[str, str]) -> str:
    normalized = description.lower()
    for pattern, category in mapping.items():
        if re.search(pattern, normalized):
            return category
    return "Uncategorized"


def normalize_transactions(df: pd.DataFrame, cfg: BuilderConfig) -> pd.DataFrame:
    required_cols = [cfg.date_col, cfg.description_col, cfg.amount_col]
    missing = [col for col in required_cols if col not in df.columns]
    if missing:
        raise ValueError(f"CSV is missing required columns: {', '.join(missing)}")

    tx = df[required_cols].copy()
    tx.columns = ["Date", "Description", "Amount"]

    tx["Date"] = pd.to_datetime(tx["Date"], errors="coerce")
    tx["Amount"] = pd.to_numeric(tx["Amount"], errors="coerce")
    tx = tx.dropna(subset=["Date", "Amount"]).sort_values("Date")

    if cfg.debit_positive:
        tx["SignedAmount"] = tx["Amount"] * -1
    else:
        tx["SignedAmount"] = tx["Amount"]

    tx["Type"] = tx["SignedAmount"].apply(lambda x: "Income" if x > 0 else "Expense")
    tx["Category"] = tx["Description"].astype(str).apply(
        lambda text: categorize(text, DEFAULT_CATEGORY_MAP)
    )
    tx["Month"] = tx["Date"].dt.to_period("M").astype(str)

    tx["Income"] = tx["SignedAmount"].apply(lambda x: x if x > 0 else 0)
    tx["Expense"] = tx["SignedAmount"].apply(lambda x: abs(x) if x < 0 else 0)

    return tx[
        [
            "Date",
            "Description",
            "Category",
            "Type",
            "Amount",
            "SignedAmount",
            "Income",
            "Expense",
            "Month",
        ]
    ]


def build_category_summary(tx: pd.DataFrame) -> pd.DataFrame:
    grouped = (
        tx.groupby("Category", dropna=False)
        .agg(
            TotalIncome=("Income", "sum"),
            TotalExpense=("Expense", "sum"),
        )
        .reset_index()
    )
    grouped["Net"] = grouped["TotalIncome"] - grouped["TotalExpense"]
    return grouped.sort_values("TotalExpense", ascending=False)


def build_monthly_summary(tx: pd.DataFrame) -> pd.DataFrame:
    grouped = (
        tx.groupby("Month")
        .agg(
            Income=("Income", "sum"),
            Expense=("Expense", "sum"),
        )
        .reset_index()
        .sort_values("Month")
    )
    grouped["Net"] = grouped["Income"] - grouped["Expense"]
    return grouped


def overwrite_sheet(workbook, sheet_name: str, rows: Iterable[Iterable], header: Optional[list] = None):
    if sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        workbook.remove(ws)
    ws = workbook.create_sheet(sheet_name)

    if header:
        ws.append(header)
    for row in rows:
        ws.append(list(row))


def write_output(template_path: Path, output_path: Path, tx: pd.DataFrame):
    workbook = load_workbook(template_path)

    category_summary = build_category_summary(tx)
    monthly_summary = build_monthly_summary(tx)

    overwrite_sheet(
        workbook,
        "Transactions",
        dataframe_to_rows(tx, index=False, header=False),
        header=list(tx.columns),
    )
    overwrite_sheet(
        workbook,
        "Category Summary",
        dataframe_to_rows(category_summary, index=False, header=False),
        header=list(category_summary.columns),
    )
    overwrite_sheet(
        workbook,
        "Monthly Summary",
        dataframe_to_rows(monthly_summary, index=False, header=False),
        header=list(monthly_summary.columns),
    )

    workbook.save(output_path)


def main() -> None:
    cfg = parse_args()
    df = pd.read_csv(cfg.csv_path)
    tx = normalize_transactions(df, cfg)
    write_output(cfg.template_path, cfg.output_path, tx)
    print(f"✅ Budget workbook created at {cfg.output_path}")


if __name__ == "__main__":
    main()
